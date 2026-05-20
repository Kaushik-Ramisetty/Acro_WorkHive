"""Mirror admin UI edits back into Employees-updated.xlsx so the spreadsheet
stays in sync with the database.

Behaviour:
- Best-effort. Any failure (file locked by Excel/OneDrive, sheet missing,
  openpyxl import error) is logged and swallowed -- the DB write that came
  before is the source of truth at runtime, and the next successful sync
  will catch the xlsx up.
- Field mapping mirrors the column layout of the `employees` sheet that
  seed.py reads. If you reorder columns in the xlsx, update HEADER_TO_COL.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

# 1-indexed column positions in the `employees` sheet (matches seed.py).
HEADER_TO_COL = {
    "id": 1,
    "employee_code": 2,
    "first_name": 4,
    "last_name": 5,
    "email": 6,
    "phone": 7,
    "designation_id": 12,
    "employment_status": 14,
    "is_deleted": 24,
    "deleted_at": 29,
    "department_id": 30,
    "location": 32,
}

# Backend root: app/utils/excel_sync.py -> Backend/
_BACKEND_ROOT = Path(__file__).resolve().parent.parent.parent
_PROJECT_ROOT = _BACKEND_ROOT.parent


def find_xlsx() -> Path | None:
    """Same discovery order as seed.find_xlsx()."""
    candidates = [
        _BACKEND_ROOT / "Employees-updated.xlsx",
        _PROJECT_ROOT / "Employees-updated.xlsx",
        _BACKEND_ROOT / "data" / "Employees-updated.xlsx",
    ]
    if _PROJECT_ROOT.exists():
        try:
            candidates.extend(sorted(_PROJECT_ROOT.glob("Employees*.xlsx")))
        except Exception:
            pass

    seen: set[Path] = set()
    for p in candidates:
        try:
            rp = p.resolve()
        except Exception:
            continue
        if rp in seen:
            continue
        seen.add(rp)
        if rp.is_file():
            return rp
    return None


def _coerce_for_excel(value: Any) -> Any:
    """openpyxl handles most Python types, but timezone-aware datetimes need
    to be stripped (Excel has no concept of tz)."""
    from datetime import datetime
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def update_employee_row(employee_id: int, updates: dict[str, Any]) -> bool:
    """Patch the row in the `employees` sheet whose first cell == employee_id.

    `updates` keys must be from HEADER_TO_COL. Unknown keys are skipped.
    Returns True if the workbook was saved, False otherwise (with a log line).
    """
    if not updates:
        return False

    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("openpyxl not installed; cannot sync to xlsx")
        return False

    xlsx = find_xlsx()
    if not xlsx:
        logger.info("Employees-updated.xlsx not found; skipping xlsx sync for id=%s", employee_id)
        return False

    try:
        wb = load_workbook(xlsx)
    except PermissionError:
        logger.warning(
            "xlsx is locked (Excel open or OneDrive mid-sync); skipping xlsx sync for id=%s",
            employee_id,
        )
        return False
    except Exception as e:
        logger.warning("Failed to open xlsx for sync (id=%s): %s", employee_id, e)
        return False

    if "employees" not in wb.sheetnames:
        logger.warning("xlsx has no 'employees' sheet; skipping sync")
        return False

    ws = wb["employees"]

    # Find the row whose id-cell matches.
    target_row = None
    for row_idx in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=HEADER_TO_COL["id"]).value
        try:
            if cell is not None and int(cell) == int(employee_id):
                target_row = row_idx
                break
        except (TypeError, ValueError):
            continue

    if target_row is None:
        logger.info("Employee id=%s not in xlsx; skipping sync", employee_id)
        return False

    applied = []
    for field, value in updates.items():
        col = HEADER_TO_COL.get(field)
        if not col:
            continue
        ws.cell(row=target_row, column=col).value = _coerce_for_excel(value)
        applied.append(field)

    if not applied:
        return False

    try:
        wb.save(xlsx)
    except PermissionError:
        logger.warning(
            "xlsx locked on save; xlsx not updated for id=%s. DB still has the change.",
            employee_id,
        )
        return False
    except Exception as e:
        logger.warning("Failed to save xlsx for id=%s: %s", employee_id, e)
        return False

    logger.info("Synced xlsx row %d (employee id=%s): %s", target_row, employee_id, applied)
    return True
